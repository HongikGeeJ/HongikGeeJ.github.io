#!/usr/bin/env python3
"""Rebuild SALES_DB_SEED from live xlsx mirroring fixed JS sales parser."""
import ctypes
import json
import math
import re
from collections import Counter
from datetime import date, datetime

import openpyxl

WB_PATH = "TeacherLogin/data/hongik-sales-live.xlsx"
OUT_PATH = "TeacherLogin/data/sales-db-seed-v3.json"

MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}


def sales_sheet_month(name):
    s = str(name or "")
    if re.search(r"table", s, re.I):
        return None
    m = re.search(
        r"(january|february|march|april|may|june|july|august|september|october|november|december)\s*2026",
        s, re.I,
    )
    if m:
        return {"year": 2026, "month": MONTHS[m.group(1).lower()]}
    m = re.search(r"(0?[1-9]|1[0-2])\s*[A-Za-z]*\s*2026", s)
    if m:
        return {"year": 2026, "month": int(m.group(1))}
    return None


def fix_serial_for_sheet(y, m, day, sheet_year, sheet_month):
    if y not in (sheet_year, sheet_year - 1, sheet_year + 1):
        return None
    if abs(m - sheet_month) <= 1 or (sheet_month == 1 and m == 12) or (sheet_month == 12 and m == 1):
        if y == 2026:
            return f"{y:04d}-{m:02d}-{day:02d}"
        return None
    if 1 <= m <= 12 and (day == sheet_month or day + 1 == sheet_month):
        try:
            datetime(sheet_year, sheet_month, m)
            return f"{sheet_year:04d}-{sheet_month:02d}-{m:02d}"
        except ValueError:
            pass
    if day <= 12 and m <= 12:
        try:
            iso = f"{sheet_year:04d}-{day:02d}-{m:02d}"
            dt = datetime(sheet_year, day, m)
            if abs(dt.month - sheet_month) <= 1:
                return iso
        except ValueError:
            pass
    if y == 2026 and 2 <= m <= 12:
        return f"{y:04d}-{m:02d}-{day:02d}"
    return None


def sales_parse_date(v, prefer_ym):
    py, pm = prefer_ym["year"], prefer_ym["month"]
    if isinstance(v, datetime):
        fixed = fix_serial_for_sheet(v.year, v.month, v.day, py, pm)
        return fixed or f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
    if isinstance(v, date):
        fixed = fix_serial_for_sheet(v.year, v.month, v.day, py, pm)
        return fixed or f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
    s = str(v or "").strip()
    if not s or s in ("-", "—"):
        return None
    m = re.match(r"^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?", s)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        yr = int(m.group(3)) if m.group(3) else py
        if yr < 100:
            yr += 2000
        cands = []
        if 1 <= b <= 12 and 1 <= a <= 31:
            cands.append((a, b))
        if 1 <= a <= 12 and 1 <= b <= 31:
            cands.append((b, a))
        best, best_score = None, 999
        for day, mon in cands:
            use_yr = py if mon == pm else yr
            try:
                datetime(use_yr, mon, day)
            except ValueError:
                continue
            score = abs(mon - pm)
            if use_yr != py:
                score += 20
            if score < best_score:
                best_score = score
                best = f"{use_yr:04d}-{mon:02d}-{day:02d}"
        return best
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if not m:
        return None
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if mo == pm and y != py:
        return f"{py:04d}-{mo:02d}-{d:02d}"
    return f"{y:04d}-{mo:02d}-{d:02d}"


def sales_parse_money(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        n = float(s)
        return n if math.isfinite(n) else None
    except ValueError:
        return None


def sales_junk_name(name):
    s = str(name or "").strip().lower()
    return (not s) or len(s) < 2 or bool(
        re.search(r"new students|existing|total|summary|students$|existin", s)
    )


def sales_name_key(name):
    return re.sub(r"[^a-z0-9ก-๙]", "", str(name or "").lower())


def sales_course_kind(course, promo, amount):
    blob = f"{course or ''} {promo or ''}".lower()
    if amount and amount <= 500 and re.search(r"book|hangul|2b book|materials", blob):
        return "book"
    if re.search(r"english|อังกฤษ", blob):
        if re.search(r"kid|เด็ก", blob):
            return "english_kids"
        if re.search(r"1:1|interview", blob):
            return "english_private"
        return "english"
    if re.search(r"korean|เกาหลี|topik", blob):
        if re.search(r"online|onine", blob):
            return "korean_online"
        if re.search(r"1:1|2:1|conversation", blob):
            return "korean_private"
        if re.search(r"night|ค่ำ|20:00|19:00|08:00\s*pm", blob):
            return "korean_night"
        return "korean"
    if re.search(r"book", blob) and amount and amount <= 500:
        return "book"
    return "other"


def hash_str(s):
    h = ctypes.c_int32(0)
    for c in str(s or ""):
        h = ctypes.c_int32(((h.value << 5) - h.value) + ord(c))
    return h.value


def main():
    wb = openpyxl.load_workbook(WB_PATH, data_only=True)
    payments = []
    seen = set()
    dup = 0
    skipped = []

    for sname in wb.sheetnames:
        ym = sales_sheet_month(sname)
        if not ym:
            skipped.append(sname)
            continue
        month_key = f"{ym['year']}-{ym['month']:02d}"
        ws = wb[sname]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        pay_date_col = -1
        paid_col = -1
        for row in rows:
            if not row:
                continue
            blob = " ".join(str(x or "") for x in row[:10]).lower()
            if "name" in blob and ("payment" in blob or "course" in blob):
                for ci, cell in enumerate(row):
                    s = re.sub(r"\s+", " ", str(cell or "").strip().lower())
                    if re.search(r"payment\s*date", s):
                        pay_date_col = ci
                    if re.search(r"^paid\b|paid\s*\+|book", s) and paid_col < 0:
                        paid_col = ci
                break

        for row in rows:
            if not row or len(row) < 4:
                continue
            blob = " ".join(str(x or "") for x in row[:4]).lower()
            if "name" in blob and "course" in blob:
                continue
            name, name_i = "", -1
            for i in (1, 0, 2):
                if name or i >= len(row):
                    continue
                s = str(row[i] or "").strip()
                if not s or len(s) < 2 or re.match(r"^\d+(\.0)?$", s):
                    continue
                if re.match(r"^(name|course|payment|no\.?|note|discount)", s, re.I):
                    continue
                if re.search(r"students payment|2026", s, re.I) and len(s) < 36:
                    continue
                if isinstance(row[i], (int, float)):
                    continue
                name, name_i = s, i
            if not name or sales_junk_name(name):
                continue
            course = ""
            for j in range(name_i + 1, min(name_i + 3, len(row))):
                cs = str(row[j] or "").strip()
                if not cs or cs in ("-", "—"):
                    continue
                if re.search(r"course|korean|english|book|topik|lesson|class|interview|online", cs, re.I) or len(cs) > 12:
                    course = re.sub(r"\s+", " ", cs)
                    break
            paid = None
            if paid_col >= 0 and paid_col < len(row):
                n = sales_parse_money(row[paid_col])
                if n is not None and n >= 0:
                    paid = n
            if paid is None:
                for j in (7, 8, 6, 5, 9):
                    if paid is not None or j >= len(row):
                        continue
                    n = sales_parse_money(row[j])
                    if n is not None and n >= 0:
                        paid = n
            if paid is None:
                for j in range(len(row) - 1, name_i, -1):
                    n = sales_parse_money(row[j])
                    if n is not None and 0 < n < 100000:
                        paid = n
                        break
            if paid is None or paid < 50:
                continue
            pay_date, promo, note = None, "", ""
            if pay_date_col >= 0 and pay_date_col < len(row):
                pay_date = sales_parse_date(row[pay_date_col], ym)
            if not pay_date:
                for j in range(name_i + 1, min(name_i + 6, len(row))):
                    if course and str(row[j] or "").strip()[:18] == course[:18]:
                        continue
                    if pay_date_col >= 0 and (j == pay_date_col + 1 or j == pay_date_col + 2):
                        continue
                    d = sales_parse_date(row[j], ym)
                    if d:
                        pay_date = d
                        break
            for cell in row:
                s = str(cell or "").strip()
                if not s or sales_parse_money(s) is not None:
                    continue
                if re.search(r"%|off|discount|deposit|book", s, re.I):
                    if not promo:
                        promo = s
                    elif not note:
                        note = s
                elif re.match(r"^\(.*\)$", s) or (len(s) > 25 and re.search(r"discount|month|promotion", s, re.I)):
                    note = s
            kind = sales_course_kind(course, f"{promo} {note}", paid)
            nk = sales_name_key(name)
            ck = re.sub(r"[^a-z0-9]", "", (course or promo or kind).lower())[:48]
            # Include paymentDate so same student/course/amount on different days stay separate.
            dedupe = f"{nk}|{ck}|{round(paid)}|{month_key}|{pay_date or ''}"
            if dedupe in seen:
                dup += 1
                continue
            seen.add(dedupe)
            h = abs(hash_str(dedupe))
            name_clean = re.sub(r"\s+", " ", name.replace("\n", " ")).strip()
            name_disp = re.sub(r"^(khun|k'|miss|nong|พี่|คุณ)\s*", "", name_clean, flags=re.I).strip()
            payments.append({
                "id": f"pay-{len(dedupe)}{format(h & 0xFFFFFFFF, 'x')[:8]}",
                "month": month_key,
                "name": name_disp,
                "nameRaw": name_clean,
                "course": course or promo or kind,
                "kind": kind,
                "paymentDate": pay_date or f"{month_key}-15",
                "promotion": promo,
                "note": note,
                "amount": round(paid * 100) / 100,
                "sheet": str(sname or "").strip(),
            })

    payments.sort(key=lambda p: (p["month"], p["paymentDate"], p["name"]))

    by_month, by_kind = {}, {}
    for p in payments:
        if p["month"] not in by_month:
            by_month[p["month"]] = {"revenue": 0, "count": 0, "students": set(), "byKind": {}}
        m = by_month[p["month"]]
        m["revenue"] += p["amount"]
        m["count"] += 1
        m["students"].add(sales_name_key(p["name"]))
        if p["kind"] not in m["byKind"]:
            m["byKind"][p["kind"]] = {"revenue": 0, "count": 0}
        m["byKind"][p["kind"]]["revenue"] += p["amount"]
        m["byKind"][p["kind"]]["count"] += 1
        if p["kind"] not in by_kind:
            by_kind[p["kind"]] = {"revenue": 0, "count": 0}
        by_kind[p["kind"]]["revenue"] += p["amount"]
        by_kind[p["kind"]]["count"] += 1

    months, prev = [], None
    for mk in sorted(by_month):
        m = by_month[mk]
        rev = round(m["revenue"] * 100) / 100
        trend = None if prev is None or prev <= 0 else round(((rev - prev) / prev) * 1000) / 10
        months.append({
            "month": mk,
            "revenue": rev,
            "payments": m["count"],
            "students": len(m["students"]),
            "trendPct": trend,
            "byKind": {
                k: {"revenue": round(v["revenue"] * 100) / 100, "count": v["count"]}
                for k, v in m["byKind"].items()
            },
        })
        prev = rev

    total = sum(p["amount"] for p in payments)
    labels = {
        "korean": "Korean", "korean_online": "Korean Online", "korean_private": "Korean 1:1/2:1",
        "korean_night": "Korean Night", "english": "English", "english_kids": "English Kids",
        "english_private": "English 1:1", "book": "Books / materials", "other": "Other",
    }
    kinds = sorted(
        [{
            "kind": k,
            "label": labels.get(k, k),
            "revenue": round(v["revenue"] * 100) / 100,
            "count": v["count"],
            "sharePct": round((v["revenue"] / total) * 1000) / 10 if total else 0,
        } for k, v in by_kind.items()],
        key=lambda x: -x["revenue"],
    )
    book = sum(p["amount"] for p in payments if p["kind"] == "book")
    seed = {
        "version": 3,
        "sourceId": "1eIHPemhSpgKWPLzeCr17trKswPxWDMzhK4t_XdH1F00",
        "sourceTitle": "Students Payment 2026",
        "importedAt": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%S"),
        "range": f"{months[0]['month']}..{months[-1]['month']}" if months else "",
        "stats": {
            "totalRevenue": round(total * 100) / 100,
            "tuitionRevenue": round((total - book) * 100) / 100,
            "bookRevenue": round(book * 100) / 100,
            "payments": len(payments),
            "uniqueStudents": len({sales_name_key(p["name"]) for p in payments}),
            "months": len(months),
            "duplicatesSkipped": dup,
        },
        "kindLabels": labels,
        "months": months,
        "kinds": kinds,
        "payments": payments,
        "skippedSheets": skipped,
    }

    sep = next(m for m in months if m["month"] == "2026-09")
    print("TOTAL", seed["stats"])
    print("SEP", sep)
    print("SEP paymentDates:")
    for p in payments:
        if p["month"] == "2026-09":
            print(f"  {p['paymentDate']} {p['amount']:8} {p['name'][:50]}")
    print("years", dict(Counter(p["paymentDate"][:4] for p in payments)))
    weird = [p for p in payments if not p["paymentDate"].startswith("2026")]
    print("non-2026", len(weird))
    for p in weird[:10]:
        print(" ", p["paymentDate"], p["month"], p["name"][:30])

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        json.dump(seed, f, ensure_ascii=False, separators=(",", ":"))
    print("wrote", OUT_PATH)


if __name__ == "__main__":
    main()
